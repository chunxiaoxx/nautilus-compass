"""Build a clean task fuel digest from batch task zip + feedback sheet."""

from __future__ import annotations

import argparse
import io
import json
import re
from collections import Counter
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from zipfile import BadZipFile, ZipFile

import openpyxl


TEXT_EXTS = {".txt", ".md", ".json", ".log", ".yaml", ".yml"}
SPREADSHEET_EXTS = {".xlsx"}
DOC_EXTS = {".docx"}


@dataclass
class TaskEntry:
    task_id: str
    domain: str
    owner: str
    title: str
    package: str
    file_count: int
    attachments: List[str]
    source_rows: List[str]
    difficulty: str = "未分类"
    feedback_issues: List[str] = None
    feedback_notes: List[str] = None
    raw_snippet: str = ""
    raw_excerpt_lines: List[str] = None

    def __post_init__(self) -> None:
        if self.feedback_issues is None:
            self.feedback_issues = []
        if self.feedback_notes is None:
            self.feedback_notes = []
        if self.raw_excerpt_lines is None:
            self.raw_excerpt_lines = []


def clean_text(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    s = re.sub(r"\s+", " ", s)
    return s.replace("\u0000", "").strip()


def normalize_owner(name: str) -> str:
    name = clean_text(name)
    name = re.sub(r"\s+", "", name)
    name = re.sub(r"[+._-]", "", name)
    return name


def similarity(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, normalize_owner(a), normalize_owner(b)).ratio()


def extract_docx_text(raw: bytes) -> str:
    try:
        with ZipFile(io.BytesIO(raw)) as zf:
            with zf.open("word/document.xml") as f:
                xml = f.read().decode("utf-8", errors="ignore")
        xml = re.sub(r"<w:p[^>]*>", "\n", xml)
        xml = re.sub(r"<[^>]+>", "", xml)
        xml = re.sub(r"\s+", " ", xml).strip()
        return xml
    except Exception:
        return ""


def extract_xlsx_text(raw: bytes) -> str:
    try:
        with io.BytesIO(raw) as bio:
            wb = openpyxl.load_workbook(bio, data_only=True, read_only=True)
            texts: List[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(min_row=1, max_row=20, max_col=20, values_only=True):
                    line = " ".join(clean_text(v) for v in row if clean_text(v))
                    if line:
                        texts.append(line)
                        if len(texts) >= 12:
                            break
                if len(texts) >= 12:
                    break
            return " ".join(texts).strip()
    except Exception:
        return ""


def safe_text_excerpt(raw: bytes, limit: int = 1200) -> str:
    return clean_text(raw[:limit].decode("utf-8", errors="ignore"))


def pick_title(raw_text: str, filename: str) -> str:
    cleaned = clean_text(raw_text)
    if not cleaned:
        return clean_text(Path(filename).stem)
    return cleaned[:120] + ("…" if len(cleaned) > 120 else "")


def parse_feedback(path: Path) -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["问题汇总"] if "问题汇总" in wb.sheetnames else wb.active
    header = [clean_text(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    idx = {name: i + 1 for i, name in enumerate(header)}

    rows = []
    for r in range(2, ws.max_row + 1):
        row = {name: clean_text(ws.cell(r, c).value) for name, c in idx.items()}
        if not any(row.values()):
            continue
        if not row.get("项目组", "").startswith("大模型训练-国曙"):
            continue

        issues = [row[f"问题{i}"] for i in range(1, 13) if row.get(f"问题{i}")]
        notes = [row[f"问题{i}"] for i in range(1, 13) if row.get(f"问题{i}")]
        attachment_raw = row.get("附件列表", "")
        attachments = [a.strip() for a in re.split(r"[\n;,]", attachment_raw) if a and a.strip()]

        rows.append(
            {
                "row_index": r,
                "project": row.get("项目组"),
                "department": row.get("部门"),
                "owner": row.get("姓名"),
                "raw_title": row.get("题目"),
                "difficulty": row.get("题型（L几）"),
                "attachments": attachments,
                "steps": row.get("做题关键步骤"),
                "issues": issues,
                "notes": notes,
            }
        )

    return rows


def parse_task_node(domain: str, node_path: str, raw: bytes) -> Optional[TaskEntry]:
    try:
        zf = ZipFile(io.BytesIO(raw))
    except BadZipFile:
        return None

    owner = node_path.split("/")[-1]
    owner = owner[:-4] if owner.endswith(".zip") else owner
    owner = owner.replace(".", "").replace("-", "")

    with zf:
        filtered = [n for n in zf.namelist() if not n.endswith("/") and not n.startswith("__MACOSX/")]
        attachments = sorted([Path(n).name for n in filtered])

        title_candidates: List[str] = []
        raw_snippet = ""
        raw_excerpt_lines: List[str] = []

        for n in filtered:
            stem = Path(n).name
            ext = Path(n).suffix.lower()

            txt = ""
            if ext in TEXT_EXTS:
                with zf.open(n) as f:
                    txt = safe_text_excerpt(f.read())
            elif ext in DOC_EXTS:
                with zf.open(n) as f:
                    txt = extract_docx_text(f.read())
            elif ext in SPREADSHEET_EXTS:
                with zf.open(n) as f:
                    txt = extract_xlsx_text(f.read())
            elif "题目" in stem or "题目数据" in stem:
                with zf.open(n) as f:
                    txt = safe_text_excerpt(f.read())

            txt = clean_text(txt)
            if txt:
                title_candidates.append(txt)
                if not raw_snippet:
                    raw_snippet = txt[:3000]
                    raw_excerpt_lines = [l for l in txt.split(" ") if l][:20]

    title = pick_title(" ".join(title_candidates[:2]), owner)

    if "连续劳务" in node_path:
        difficulty = "L2流程型"
    elif "客服" in node_path:
        difficulty = "L2流程型"
    elif "市场" in node_path:
        difficulty = "L2类型"
    elif any(k in node_path for k in ["电销", "销售", "人资"]):
        difficulty = "L2类型"
    else:
        difficulty = "偏难"

    return TaskEntry(
        task_id=node_path,
        domain=domain,
        owner=owner,
        title=title or "未命名题目",
        package=Path(node_path).name,
        file_count=len(filtered),
        attachments=attachments,
        source_rows=[node_path],
        difficulty=difficulty,
        raw_snippet=raw_snippet,
        raw_excerpt_lines=raw_excerpt_lines,
    )


def extract_tasks_from_archive(outer_zip: Path) -> List[TaskEntry]:
    tasks: List[TaskEntry] = []
    with ZipFile(outer_zip) as zf:
        for n in zf.namelist():
            if n.endswith("/"):
                continue
            if "._" in Path(n).name:
                continue
            if not n.lower().endswith(".zip"):
                continue
            if "__MACOSX/" in n:
                continue

            segments = [p for p in Path(n).parts if p]
            if len(segments) < 3:
                continue
            domain = segments[1]
            entry = parse_task_node(domain, n, zf.read(n))
            if entry:
                tasks.append(entry)
    return tasks


def build_links(tasks: List[TaskEntry], feedback_rows: List[Dict[str, Any]]) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    result = []
    linked = set()

    for task in tasks:
        best = None
        best_score = 0.0
        for i, row in enumerate(feedback_rows):
            for key in ["owner", "project", "department"]:
                s = similarity(task.owner, row.get(key, ""))
                if s > best_score:
                    best_score = s
                    best = i
            if row.get("raw_title") and clean_text(row.get("raw_title")) in task.title:
                best_score = 1.0
                best = i

        if best is not None and best_score >= 0.3:
            row = feedback_rows[best]
            linked.add(best)
            task.feedback_issues = row["issues"]
            task.feedback_notes = row["notes"]
            if row.get("difficulty"):
                task.difficulty = row["difficulty"]
            task.title = task.title or row.get("raw_title") or task.title
            task.source_rows.append(f"feedback:{row['row_index']}")
            task.raw_snippet = row.get("steps") or task.raw_snippet
            task.raw_excerpt_lines = row.get("issues")[:6]

        result.append(asdict(task))

    for i, row in enumerate(feedback_rows):
        if i in linked:
            continue
        owner = row.get("owner", "")
        if not owner:
            continue

        dep = row.get("department", "")
        if "电销" in dep:
            domain = "电销/销售"
        elif "客服" in dep:
            domain = "客服"
        elif "销售" in dep:
            domain = "销售"
        elif "市场" in dep:
            domain = "市场"
        elif "人力" in dep:
            domain = "人力资源"
        elif "连续劳务" in dep:
            domain = "连续劳务"
        elif dep:
            domain = dep
        else:
            domain = "未分类"

        result.append(
            {
                "task_id": f"feedback-only:{owner}:{i}",
                "domain": domain,
                "owner": owner,
                "title": row.get("raw_title", "未命名题目"),
                "package": "",
                "file_count": len(row.get("attachments", [])),
                "attachments": row.get("attachments", []),
                "source_rows": [f"feedback:{row['row_index']}"],
                "difficulty": row.get("difficulty") or "未分类",
                "feedback_issues": row.get("issues", []),
                "feedback_notes": row.get("notes", []),
                "raw_snippet": row.get("steps", ""),
                "raw_excerpt_lines": row.get("issues", [])[:6],
            }
        )

    by_domain = Counter([t["domain"] for t in result])
    by_difficulty = Counter([t["difficulty"] or "未分类" for t in result])
    issue_count = sum(1 for t in result for _ in t.get("feedback_issues", []))

    issue_top = Counter()
    for t in result:
        for issue in t.get("feedback_issues", []):
            tag = re.split(r"[:：]", issue)[0].strip()
            if tag:
                if len(tag) > 20:
                    tag = tag[:20]
                issue_top[tag] += 1

    stats = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "total_tasks": len(result),
        "source_counts": dict(by_domain),
        "difficulty_counts": dict(by_difficulty),
        "issue_total": issue_count,
        "top_issue_tags": [f"{k}:{v}" for k, v in issue_top.most_common(10)],
        "failure_tags_top": [{"tag": k, "count": v} for k, v in issue_top.most_common(10)],
    }

    return result, stats


def build_digest(outer_zip: Path, feedback_xlsx: Path, output: Path) -> None:
    feedback_rows = parse_feedback(feedback_xlsx)
    tasks = extract_tasks_from_archive(outer_zip)
    task_payload, stats = build_links(tasks, feedback_rows)

    task_payload.sort(key=lambda x: (x["domain"], x["owner"]))

    digest = {
        "meta": stats,
        "tasks": task_payload,
        "links": {
            "outer_zip": str(outer_zip),
            "feedback_xlsx": str(feedback_xlsx),
        },
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(digest, ensure_ascii=False, indent=2), encoding="utf-8")


def parse_args():
    parser = argparse.ArgumentParser(description="Build embodied task fuel digest.")
    parser.add_argument("--outer-zip", required=True)
    parser.add_argument("--feedback-xlsx", required=True)
    parser.add_argument("--out", required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_digest(Path(args.outer_zip), Path(args.feedback_xlsx), Path(args.out))


if __name__ == "__main__":
    main()
